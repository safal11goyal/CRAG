import os
import time
from typing import List, Tuple
from pypdf import PdfReader
from docx import Document as DocxReader
from langchain_core.documents import Document
from langchain_core.embeddings import Embeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.vectorstores import FAISS


def load_documents(documents_dir: str) -> List[Document]:
    """Load PDF, TXT, and DOCX files from the given directory."""
    documents: List[Document] = []

    if not os.path.exists(documents_dir):
        os.makedirs(documents_dir, exist_ok=True)
        return documents

    for filename in sorted(os.listdir(documents_dir)):
        filepath = os.path.join(documents_dir, filename)
        if os.path.isdir(filepath):
            continue

        ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""

        if ext == "pdf":
            try:
                reader = PdfReader(filepath)
                for page_idx, page in enumerate(reader.pages):
                    text = page.extract_text()
                    if text and text.strip():
                        documents.append(
                            Document(
                                page_content=text,
                                metadata={
                                    "source": filename,
                                    "page": page_idx + 1,  # 1-indexed
                                },
                            )
                        )
            except Exception as e:
                print(f"[WARN] Failed to load PDF {filename}: {e}", flush=True)

        elif ext == "txt":
            try:
                with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
                    text = f.read()
                if text.strip():
                    documents.append(
                        Document(
                            page_content=text,
                            metadata={"source": filename},
                        )
                    )
            except Exception as e:
                print(f"[WARN] Failed to load TXT {filename}: {e}", flush=True)

        elif ext == "docx":
            try:
                doc = DocxReader(filepath)
                text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
                if text.strip():
                    documents.append(
                        Document(
                            page_content=text,
                            metadata={"source": filename},
                        )
                    )
            except Exception as e:
                print(f"[WARN] Failed to load DOCX {filename}: {e}", flush=True)

    return documents


def chunk_documents(
    documents: List[Document], chunk_size: int = 1000, chunk_overlap: int = 200
) -> List[Document]:
    """Split documents into chunks while preserving metadata."""
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        separators=["\n\n", "\n", " ", ""],
    )
    return splitter.split_documents(documents)


def load_or_create_vectorstore(
    documents_dir: str,
    vectorstore_dir: str,
    embeddings: Embeddings,
    chunk_size: int = 1000,
    chunk_overlap: int = 200,
    batch_size: int = 30,
) -> Tuple[FAISS, List[Document]]:
    """Load existing FAISS index or create and persist a new one in rate-limited batches with retries."""
    raw_docs = load_documents(documents_dir)
    chunks = chunk_documents(raw_docs, chunk_size=chunk_size, chunk_overlap=chunk_overlap)

    index_file = os.path.join(vectorstore_dir, "index.faiss")
    if os.path.exists(index_file):
        print(f"[INFO] Loading existing FAISS index from '{vectorstore_dir}'...", flush=True)
        vectorstore = FAISS.load_local(
            vectorstore_dir, embeddings, allow_dangerous_deserialization=True
        )
        return vectorstore, chunks

    print(f"[INFO] Creating new FAISS index from {len(chunks)} chunks in batches of {batch_size}...", flush=True)
    if not chunks:
        raise ValueError(
            f"No valid documents found in '{documents_dir}'. Please add PDF/TXT/DOCX files."
        )

    os.makedirs(vectorstore_dir, exist_ok=True)
    vectorstore = None
    total_batches = (len(chunks) + batch_size - 1) // batch_size

    for idx, i in enumerate(range(0, len(chunks), batch_size), start=1):
        batch = chunks[i : i + batch_size]
        print(f"[INFO] Embedding batch {idx}/{total_batches} ({len(batch)} chunks)...", flush=True)

        for attempt in range(8):
            try:
                if vectorstore is None:
                    vectorstore = FAISS.from_documents(batch, embeddings)
                else:
                    vectorstore.add_documents(batch)
                break
            except Exception as e:
                err_msg = str(e).lower()
                if "429" in err_msg or "rate limit" in err_msg or "too many requests" in err_msg or "resource_exhausted" in err_msg or "trial token" in err_msg:
                    wait_time = (attempt + 1) * 10
                    print(f"[WARN] Rate limit reached. Waiting {wait_time}s before retrying batch {idx}/{total_batches} (attempt {attempt + 1}/8)...", flush=True)
                    time.sleep(wait_time)
                else:
                    raise e

        # Courtesy pause between batches to stay under Cohere's 100k TPM trial limit
        time.sleep(3)

    if vectorstore is not None:
        vectorstore.save_local(vectorstore_dir)
        print(f"[INFO] FAISS index successfully created and saved to '{vectorstore_dir}'.", flush=True)

    return vectorstore, chunks
